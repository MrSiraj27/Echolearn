from pptx import Presentation

from app.documents.parsers.tables import rows_to_table_block
from app.documents.parsers.types import ParsedPage


def parse_pptx(file_path: str) -> list[ParsedPage]:
    prs = Presentation(file_path)
    pages: list[ParsedPage] = []

    for i, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        tables = []
        table_idx = 0

        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs).strip()
                    if text:
                        parts.append(text)

            if shape.has_table:
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                block = rows_to_table_block(rows, table_idx)
                if block:
                    tables.append(block)
                    table_idx += 1

        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"[Speaker notes: {notes}]")

        pages.append(ParsedPage(page_number=i, text="\n".join(parts), tables=tables))

    return pages
