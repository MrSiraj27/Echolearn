from docx import Document as DocxDocument

from app.documents.parsers.tables import rows_to_table_block
from app.documents.parsers.types import ParsedPage


def parse_docx(file_path: str) -> list[ParsedPage]:
    doc = DocxDocument(file_path)

    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())

    tables = []
    for idx, table in enumerate(doc.tables):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        block = rows_to_table_block(rows, idx)
        if block:
            tables.append(block)

    return [ParsedPage(page_number=1, text=text, tables=tables)]
